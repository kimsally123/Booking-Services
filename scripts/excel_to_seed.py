#!/usr/bin/env python3
"""แปลงไฟล์ Excel ตั้งค่าร้าน → supabase/seed.sql

ใช้:  python3 scripts/excel_to_seed.py data/BaanMax-booking-schema.xlsx > supabase/seed.sql
      --setup-only  เอาเฉพาะแท็บตั้งค่าร้าน (ใช้ตอนขึ้นระบบจริง ไม่เอาลูกค้า/การจองสมมติ)
      --no-phase2   ไม่ใส่ข้อมูลเฟส 2

อ่านแถว 2 เป็นชื่อคอลัมน์ ข้อมูลเริ่มแถว 3 · ตรวจช่องบังคับและรหัสอ้างอิงก่อน
ถ้าเจอข้อผิดพลาด พิมพ์ชื่อแท็บ + แถวให้แก้ แล้วจบโดยไม่สร้าง SQL
"""
import sys, re, datetime as dt
from openpyxl import load_workbook

# แท็บ → ตาราง (ลำดับ import)
SHEETS = [
    ('ตั้งค่าร้าน', 'settings', 'setup'), ('ธุรกิจ', 'businesses', 'setup'), ('ห้องและบ้าน', 'units', 'setup'),
    ('ขนาดน้อง', 'weight_tiers', 'setup'), ('บริการ', 'services', 'setup'), ('ราคาตามขนาด', 'service_prices', 'setup'),
    ('รูป', 'unit_photos', 'setup'), ('รอบสระ', 'pool_sessions', 'setup'), ('พนักงาน', 'staff', 'setup'),
    ('วันปิด', 'closures', 'setup'), ('ขายข้ามธุรกิจ', 'cross_sells', 'setup'),
    ('ลูกค้า', 'customers', 'data'), ('น้อง', 'pets', 'data'), ('วัคซีน', 'pet_vaccines', 'data'),
    ('ทริป', 'orders', 'data'), ('การจอง', 'bookings', 'data'), ('การชำระเงิน', 'payments', 'data'),
    ('บัตรสระ', 'passes', 'phase2'), ('โค้ดส่วนลด', 'coupons', 'phase2'),
]
DAYS = {'จ': 1, 'อ': 2, 'พ': 3, 'พฤ': 4, 'ศ': 5, 'ส': 6, 'อา': 7}
TZ = '+07'
# คอลัมน์ชนิดพิเศษ
LIST_TEXT = {('units', 'amenities'), ('staff', 'businesses'), ('pets', 'tags')}
LIST_DAYS = {('pool_sessions', 'weekdays'), ('staff', 'work_days'), ('businesses', 'closed_weekdays')}
DATETIME = {'pdpa_at', 'hold_until', 'created_at', 'start_at', 'end_at', 'reviewed_at'}
BOOL = {'active', 'need_vaccine', 'is_cover', 'is_groomer'}
SYSTEM_ONLY = {('businesses', 'calendar_id')}   # ไม่ import (ระบบเติมตอนเชื่อมต่อ)

errors = []
def err(sheet, row, msg): errors.append(f'[{sheet}] แถว {row}: {msg}')

FORMULA_WB = None   # ไฟล์ที่บันทึกจากบางโปรแกรมเก็บ TRUE/FALSE เป็นสูตร =TRUE() โดยไม่มีค่าที่คำนวณไว้

def cell(ws, r, c):
    v = ws.cell(row=r, column=c).value
    if v is None and FORMULA_WB is not None:
        f = FORMULA_WB[ws.title].cell(row=r, column=c).value
        if isinstance(f, str) and f.upper() in ('=TRUE()', '=TRUE'): return True
        if isinstance(f, str) and f.upper() in ('=FALSE()', '=FALSE'): return False
    return v

def read(wb, sheet):
    if sheet not in wb.sheetnames: return [], []
    ws = wb[sheet]
    cols = [c.value for c in ws[2]]
    n = 0
    while n < len(cols) and cols[n]: n += 1
    cols = [str(c).strip() for c in cols[:n]]
    req = ['*' in str(ws.cell(row=1, column=j + 1).value or '') for j in range(n)]
    rows = []
    for r in range(3, ws.max_row + 1):
        vals = [cell(ws, r, j + 1) for j in range(n)]
        if all(v is None or str(v).strip() == '' for v in vals): continue
        d = {}
        for c, v in zip(cols, vals):
            if isinstance(v, str): v = v.strip()
            d[c] = None if v in (None, '') else v
        for c, q in zip(cols, req):
            if q and d[c] is None: err(sheet, r, f'ช่อง {c} ต้องกรอก')
        rows.append((r, d))
    return cols, rows

def q(v):
    if v is None: return 'null'
    if isinstance(v, bool): return 'true' if v else 'false'
    if isinstance(v, (int, float)): return repr(v)
    if isinstance(v, tuple) and v[0] == 'json':
        import json as _j
        return "'" + _j.dumps(v[1], ensure_ascii=False).replace("'", "''") + "'::jsonb"
    if isinstance(v, list): return 'array[' + ','.join(q(x) for x in v) + ']' + ('::smallint[]' if v and isinstance(v[0], int) else '::text[]') if v else "'{}'"
    return "'" + str(v).replace("'", "''") + "'"

def as_bool(v, sheet, r, c):
    if v is None: return None
    if isinstance(v, bool): return v
    s = str(v).upper()
    if s in ('TRUE', 'ใช่', '1'): return True
    if s in ('FALSE', 'ไม่ใช่', '0'): return False
    err(sheet, r, f'{c} ต้องเป็น TRUE/FALSE (ได้ {v})')

def as_list(v): return [x.strip() for x in str(v).split(',') if x.strip()] if v is not None else []

def as_days(v, sheet, r, c):
    out = []
    for x in as_list(v):
        if x not in DAYS: err(sheet, r, f'{c}: ไม่รู้จักวัน "{x}" (ใช้ จ,อ,พ,พฤ,ศ,ส,อา)')
        else: out.append(DAYS[x])
    return out or None

def as_dt(v, sheet, r, c):
    if v is None: return None
    if isinstance(v, dt.datetime): return v.strftime('%Y-%m-%d %H:%M') + TZ
    s = str(v)
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}( \d{2}:\d{2})?', s): err(sheet, r, f'{c} ต้องเป็น YYYY-MM-DD HH:MM (ได้ {s})'); return None
    return s + TZ

def as_date(v, sheet, r, c):
    if v is None: return None
    if isinstance(v, (dt.date, dt.datetime)): return v.strftime('%Y-%m-%d')
    s = str(v)
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', s): err(sheet, r, f'{c} ต้องเป็น YYYY-MM-DD (ได้ {s})'); return None
    return s

def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    flags = {a for a in sys.argv[1:] if a.startswith('--')}
    if not args: print(__doc__, file=sys.stderr); sys.exit(2)
    global FORMULA_WB
    wb = load_workbook(args[0], data_only=True)
    FORMULA_WB = load_workbook(args[0])
    keep = {'setup'} | (set() if '--setup-only' in flags else {'data'}) | (set() if '--no-phase2' in flags or '--setup-only' in flags else {'phase2'})

    data = {}
    for sheet, table, grp in SHEETS:
        if grp in keep: data[table] = (sheet,) + read(wb, sheet)
    codes = {t: {d.get('code') for _, d in rows} for t, (_, _, rows) in data.items()}
    session_codes, unit_codes, service_codes = codes.get('pool_sessions', set()), codes.get('units', set()), codes.get('services', set())

    def ref(table, sheet, r, col, v):
        if v is not None and v not in codes.get(table, set()):
            err(sheet, r, f'{col} = "{v}" ไม่มีในแท็บ {next(s for s, t, _ in SHEETS if t == table)}')

    # ตรวจรหัสซ้ำ
    for t, (sheet, cols, rows) in data.items():
        if 'code' in cols:
            seen = {}
            for r, d in rows:
                if d['code'] in seen: err(sheet, r, f'รหัส {d["code"]} ซ้ำกับแถว {seen[d["code"]]}')
                seen[d['code']] = r

    out = []
    pets_link = []
    for sheet, table, grp in SHEETS:
        if table not in data: continue
        _, cols, rows = data[table]
        for r, d in rows:
            rec = {}
            for c in cols:
                if (table, c) in SYSTEM_ONLY: continue
                v = d[c]
                if (table, c) in LIST_TEXT: v = as_list(v)
                elif (table, c) in LIST_DAYS: v = as_days(v, sheet, r, c)
                elif c in BOOL: v = as_bool(v, sheet, r, c)
                elif c in DATETIME: v = as_dt(v, sheet, r, c)
                elif c in ('date_from', 'date_to', 'expires_on', 'valid_until'): v = as_date(v, sheet, r, c)
                elif c in ('phone', 'code', 'customer_code') and v is not None: v = str(v)
                elif c in ('start', 'end', 'work_start', 'work_end') and v is not None: v = str(v)[:5]
                rec[c] = v
            # ---- แปลงตามตาราง ----
            if table == 'businesses':
                # 7:100,3:50 → [{"days":7,"refund":100},...]
                raw = rec.pop('cancel_rules', None)
                rules = []
                for part in as_list(raw):
                    m = re.fullmatch(r'(\d+)\s*:\s*(\d+)', part)
                    if not m: err(sheet, r, f'cancel_rules "{part}" ต้องเป็น วัน:คืน% เช่น 7:100'); continue
                    rules.append({'days': int(m.group(1)), 'refund': int(m.group(2))})
                rec['cancel_rules'] = ('json', sorted(rules, key=lambda x: -x['days']))
                if rec.get('closed_weekdays') is None: rec['closed_weekdays'] = []
            if table == 'pool_sessions':
                rec['start_time'] = rec.pop('start'); rec['end_time'] = rec.pop('end')
                ref('units', sheet, r, 'unit_code', rec['unit_code'])
            if table in ('units', 'services', 'closures'):
                ref('businesses', sheet, r, 'business_code', rec.get('business_code'))
            if table == 'service_prices':
                ref('services', sheet, r, 'service_code', rec['service_code']); ref('weight_tiers', sheet, r, 'tier_code', rec['tier_code'])
            if table == 'unit_photos':
                o = rec.pop('owner_code')
                if o in unit_codes: rec['unit_code'] = o
                elif o in service_codes: rec['service_code'] = o
                else: err(sheet, r, f'owner_code "{o}" ไม่มีทั้งในแท็บห้องและบ้าน และแท็บบริการ')
            if table == 'closures':
                u = rec.pop('unit_code')
                if u in session_codes: rec['session_code'] = u
                elif u is not None:
                    rec['unit_code'] = u; ref('units', sheet, r, 'unit_code', u)
            if table == 'cross_sells':
                ref('businesses', sheet, r, 'from_business', rec['from_business']); ref('businesses', sheet, r, 'offer_business', rec['offer_business'])
            if table == 'staff':
                for b in rec['businesses']:
                    if b != 'all': ref('businesses', sheet, r, 'businesses', b)
                rec['invite_status'] = rec.get('invite_status') or 'pending'
                rec.pop('line_user_id', None)   # ผูกตอนกดลิงก์เชิญเท่านั้น
            if table == 'pets': ref('customers', sheet, r, 'customer_code', rec['customer_code'])
            if table == 'pet_vaccines':
                ref('pets', sheet, r, 'pet_code', rec['pet_code']); ref('staff', sheet, r, 'reviewed_by', rec.get('reviewed_by'))
            if table == 'orders': ref('customers', sheet, r, 'customer_code', rec['customer_code'])
            if table == 'payments':
                ref('orders', sheet, r, 'order_code', rec['order_code']); ref('staff', sheet, r, 'reviewed_by', rec.get('reviewed_by'))
            if table == 'bookings':
                ref('orders', sheet, r, 'order_code', rec['order_code']); ref('businesses', sheet, r, 'business_code', rec['business_code'])
                u = rec.pop('unit_code')
                if u in session_codes: rec['session_code'] = u
                elif u is not None: rec['unit_code'] = u; ref('units', sheet, r, 'unit_code', u)
                sv = as_list(rec.pop('service_code'))
                if sv:
                    rec['service_code'] = sv[0]; rec['addon_codes'] = sv[1:]
                    for s in sv: ref('services', sheet, r, 'service_code', s)
                ref('staff', sheet, r, 'staff_code', rec.get('staff_code'))
                for p in as_list(rec.pop('pet_codes')):
                    ref('pets', sheet, r, 'pet_codes', p); pets_link.append((rec['code'], p))
            if table == 'passes': ref('customers', sheet, r, 'customer_code', rec['customer_code'])
            if table == 'coupons': ref('businesses', sheet, r, 'business_code', rec.get('business_code'))
            rec = {k: v for k, v in rec.items() if v is not None}
            out.append(f"insert into {table} ({', '.join(rec)}) values ({', '.join(q(v) for v in rec.values())});")
            if table == 'bookings':
                out += [f"insert into booking_pets (booking_code, pet_code) values ({q(b)}, {q(p)});" for b, p in pets_link]
                pets_link = []

    if errors:
        print('พบข้อผิดพลาด แก้ในไฟล์ Excel แล้วรันใหม่:', file=sys.stderr)
        for e in errors: print('  ' + e, file=sys.stderr)
        sys.exit(1)

    print('-- สร้างจาก ' + args[0] + ' · ' + dt.datetime.now().strftime('%Y-%m-%d %H:%M') + ' · ห้ามแก้ไฟล์นี้ตรง ๆ ให้แก้ Excel แล้วรันสคริปต์ใหม่')
    print('begin;')
    print('\n'.join(out))
    # ตั้งตัวนับรหัสให้ต่อจากข้อมูลที่ import (กันรหัสชน)
    print("""insert into code_counters (prefix, day, n)
select split_part(code, '-', 1), split_part(code, '-', 2), max(split_part(code, '-', 3)::int)
from (select code from bookings union all select code from orders) x
where code ~ '^[A-Z]+-[0-9]{6}-[0-9]+$' group by 1, 2
on conflict (prefix, day) do update set n = greatest(code_counters.n, excluded.n);
insert into code_counters (prefix, day, n)
select 'PAY', '', max(split_part(code, '-', 2)::int) from payments where code ~ '^PAY-[0-9]+$' having count(*) > 0
on conflict (prefix, day) do update set n = greatest(code_counters.n, excluded.n);""")
    # Villa: ยอดคงเหลือก่อนเช็กอิน 7 วัน ถ้าใน Excel ยังไม่มีคอลัมน์ balance_days
    print("update businesses set balance_days = 7 where code = 'villa' and balance_days is null;")
    print('commit;')
    n = {t: len(rows) for t, (_, _, rows) in data.items()}
    print('สร้าง seed แล้ว: ' + ' · '.join(f'{t} {c}' for t, c in n.items()), file=sys.stderr)

if __name__ == '__main__':
    main()
